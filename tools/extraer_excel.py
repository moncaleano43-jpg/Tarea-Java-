import openpyxl,json,datetime,sys
p='entrada/01. CONTROL  PROCESO CAVAS 2026.xlsm'
wv=openpyxl.load_workbook(p,read_only=True,data_only=True)
def iso(v):
    if isinstance(v,datetime.datetime): return v.strftime('%Y-%m-%dT%H:%M')
    return None
def n(v,zero=False):
    if isinstance(v,bool) or v is None: return None
    if isinstance(v,(int,float)):
        if v==0 and not zero: return None
        return round(float(v),4) if not float(v).is_integer() else int(v)
    return None
def s(v):
    if v is None: return None
    v=str(v).strip()
    return v if v and not v.startswith('#') else None
# ---- fermentación
ws=wv['B.D FERMENTACIÓN']
rows=list(ws.iter_rows(min_row=13,max_row=5676,max_col=260,values_only=True))
ll={}
for r in rows:
    if not r[1] or not iso(r[60]) or not iso(r[60]).startswith('2026'): continue
    cocs=[]
    for k in range(5):
        b=8+7*k
        if r[b+1] is None and r[b+3] is None: continue
        cocs.append({'ok':s(r[b]),'coc':s(r[b+1]),'temp':n(r[b+2]),'vol':n(r[b+3]),'eo':n(r[b+4]),'aire':n(r[b+5]),'lev':n(r[b+6])})
    ms=[]
    c=81
    while c+4<len(r):
        t=iso(r[c])
        if t: ms.append({'t':t,'h':n(r[c+1],True),'ext':n(r[c+2],True),'att':n(r[c+3]),'ph':n(r[c+4])})
        c+=5
    L={'lote':s(r[1]),'tq':n(r[2]),'marca':s(r[3]),'cocs':cocs,
       'aireGHl':[n(r[43]),n(r[44]),n(r[45]),n(r[46])],'pnc':s(r[49]),
       'colector':n(r[51]),'levSembrada':s(r[52]),'utqOrigen':n(r[53]),'fechaRecup':iso(r[54]),'gen':n(r[55]),
       'cons':n(r[56]),'viab':n(r[57]),'ph':n(r[58]),'tempSiembra':n(r[59]),
       'inicio':iso(r[60]),'fin':iso(r[61]),'presion':n(r[62]),'nCoc':n(r[63]),'tLlenado':n(r[64]),'vol':n(r[65]),
       'o2':n(r[66]),'eo':n(r[67]),'eLim':n(r[68]),'phMosto':n(r[69]),'amargo':n(r[70]),'color':n(r[71]),
       'levTotal':n(r[74]),'factorDosis':n(r[75]),'fechaRec3h':iso(r[76]),'rec3h':n(r[77]),'kpiHig':s(r[78]),
       'fechaMuestra72':iso(r[79]),'ea72':n(r[80]),'muestras':ms}
    ll[L['lote']]=L
print('llenados',len(ll),file=sys.stderr)
# ---- maduración
ws=wv['B.D MADURACIÓN']
rows=list(ws.iter_rows(min_row=13,max_row=5671,max_col=100,values_only=True))
tr={}
for r in rows:
    if not r[1] or not iso(r[10]) or not iso(r[10]).startswith('2026'): continue
    M={'lote':s(r[1]),'tq':n(r[2]),'marca':s(r[3]),
       'fvs':[{'lote':s(r[8]),'tq':n(r[9]),'ini':iso(r[10]),'fin':iso(r[11]),'silicaLote':s(r[12]),'silicaKg':n(r[13]),'temp':n(r[14]),
               'volAj':n(r[15]),'volCont':n(r[16]),'resp':s(r[17]),'eo':n(r[20]),'ph':n(r[21]),'colorCoc':n(r[22]),'colorFv':n(r[23]),
               'amargoCoc':n(r[24]),'amargoFv':n(r[25]),'tFV':n(r[28])}],
       'ini':iso(r[50]),'fin':iso(r[51]),'tLlenado':n(r[52]),'vol':n(r[53]),'eoPond':n(r[57]),'phPond':n(r[58]),
       'fechaFinal':iso(r[64]),'eoFinal':n(r[65]),'eapp':n(r[66]),'ereal':n(r[67]),'alcPeso':n(r[68]),'alcVol':n(r[69]),
       'rdf':n(r[70]),'colorFinal':n(r[71]),'phFin':n(r[72]),'amargoFinal':n(r[73]),'recuento':n(r[74]),'o2d':n(r[75]),'co2':n(r[76]),'turb':n(r[77]),
       'resp':s(r[78]),'sensorial':s(r[79]),'degustador':s(r[80])}
    # segundo FV (cols 29-49)
    if r[29] or iso(r[31]):
        M['fvs'].append({'lote':s(r[29]),'tq':n(r[30]),'ini':iso(r[31]),'fin':iso(r[32]),'silicaLote':s(r[33]),'silicaKg':n(r[34]),'temp':n(r[35]),
               'volAj':n(r[36]),'volCont':n(r[37]),'resp':s(r[38]),'eo':n(r[41]),'ph':n(r[42])})
    tr[M['lote']]=M
print('trasiegos',len(tr),file=sys.stderr)
out={'llenados':ll,'trasiegos':tr}
json.dump(out,open('src/seed_cava.json','w'),ensure_ascii=False,separators=(',',':'))
import os; print(os.path.getsize('src/seed_cava.json'))
k=next(iter(ll)); print(json.dumps(ll[k],ensure_ascii=False)[:1500])
k=next(iter(tr)); print(json.dumps(tr[k],ensure_ascii=False)[:1200])

# ---------- limpieza, estado de tanques, espec ----------
def strip(o):
    if isinstance(o,dict): return {k:strip(v) for k,v in o.items() if v is not None and v!=[] and v!={}}
    if isinstance(o,list): return [strip(x) for x in o]
    return o
def rnd(o):
    if isinstance(o,dict): return {k:rnd(v) for k,v in o.items()}
    if isinstance(o,list): return [rnd(x) for x in o]
    if isinstance(o,float): return round(o,3)
    return o
llenados={}
for L in ll.values():
    for m in L['muestras']: m.pop('att',None)
    llenados[L['lote']+'-'+L['inicio'][:4]]=L
trasiegos={}
for M in tr.values():
    trasiegos[M['lote']+'-'+M['ini'][:4]]=M
# marcar FV vaciados por trasiego (usa lote origen)
for tid,M in trasiegos.items():
    for fv in M['fvs']:
        if not fv.get('lote'): continue
        for k in (fv['lote']+'-2026',fv['lote']+'-2025'):
            if k in llenados and not llenados[k].get('vaciado') and llenados[k]['fin']<= (fv.get('ini') or '9'):
                llenados[k]['vaciado']=fv['ini']; llenados[k]['trasiego']=tid; break
# un tanque queda vacío cuando se registra un uso posterior en el mismo UTQ
inicios=[(x['inicio'],x['tq'],'F',k) for k,x in llenados.items()]+[(x['ini'],x['tq'],'M',k) for k,x in trasiegos.items()]
for k,x in trasiegos.items():
    sig=sorted(t for t,tq,_,kk in inicios if tq==x['tq'] and t>x['ini'])
    if sig: x['vaciado']=sig[0]
for k,x in llenados.items():
    if not x.get('vaciado'):
        sig=sorted(t for t,tq,_,kk in inicios if tq==x['tq'] and t>x['inicio'])
        if sig: x['vaciado']=sig[0]
# INVENTARIO (estado actual de los 22 tanques)
ws=wv['INVENTARIO']
cava={}
for r in ws.iter_rows(min_row=6,max_row=37,min_col=2,max_col=12,values_only=True):
    tq=n(r[0])
    if not tq: continue
    cava[str(tq)]={'tq':tq,'op':s(r[1]),'fechaLlenado':iso(r[2]),'marca':s(r[3]),'temp':n(r[5]),'levadura':s(r[6]),'fechaRecoleccion':iso(r[7]),'volReal':n(r[9])}
# etiquetar lote actual
for tq,c in cava.items():
    if not c.get('op'): continue
    src=llenados if c['op']=='F' else trasiegos
    best=None
    for k,x in src.items():
        if x['tq']==c['tq'] and (x.get('fin') or x.get('ini'))==(c['fechaLlenado']): best=k
    if not best:
        cands=[(x.get('inicio') or x.get('ini'),k) for k,x in src.items() if x['tq']==c['tq']]
        if cands: best=max(cands)[1]
    c['ref']=best
# todo registro que no sea el actual de algún tanque ya está vacío (evita reactivar tanques viejos al sincronizar)
refs={c.get('ref') for c in cava.values() if c.get('ref')}
for k,x in llenados.items():
    if k not in refs and not x.get('vaciado'): x['vaciado']=x.get('fin') or x['inicio']
for k,x in trasiegos.items():
    if k not in refs and not x.get('vaciado'): x['vaciado']=x.get('fin') or x['ini']
# espec por marca
ws=wv['ESPECIFICACIONES MARCA']
rows=list(ws.iter_rows(min_row=10,max_row=29,min_col=1,max_col=17,values_only=True))
marcas=[s(x) for x in rows[0][3:9]]
espec={'marcas':marcas,'sup':{},'inf':{}}
for r in rows[1:]:
    nom=s(r[2])
    if not nom: continue
    espec['sup'][nom]={m:n(v,True) for m,v in zip(marcas,r[3:9])}
    nomi=s(r[10]) or nom
    espec['inf'][nomi]={m:n(v,True) for m,v in zip(marcas,r[11:17])}
rows=list(ws.iter_rows(min_row=53,max_row=55,min_col=3,max_col=9,values_only=True))
espec['descanso']={m:{'ext':n(a,True),'rata':n(b,True)} for m,a,b in zip(marcas,rows[1][1:],rows[2][1:])}
out={'llenados':strip(rnd(llenados)),'trasiegos':strip(rnd(trasiegos)),'cava':strip(cava),'espec':espec}
json.dump(out,open('src/seed_cava.json','w'),ensure_ascii=False,separators=(',',':'))
print('final',os.path.getsize('src/seed_cava.json'))
print(json.dumps(cava['14'] if '14' in cava else list(cava.items())[:2],ensure_ascii=False))
print(json.dumps(espec,ensure_ascii=False)[:900])
print(sum(1 for x in llenados.values() if x.get('vaciado')),'vaciados of',len(llenados))
